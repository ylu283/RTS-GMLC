"""Build bo_1013.pptx — BO campaign + ERCOT v2 status deck (PI meeting), 2026-10-13.

Run:  /opt/anaconda3/bin/python make_bo_1013_deck.py      (from any directory)
Out:  bo_1013.pptx, SLIDE_TEXT_1013.md, FIGURES_USED_1013.md   (next to this script)

Content source of truth: OUTLINE_1013.md in this folder. This script only
renders it. Do not add numbers here that are not in OUTLINE_1013.md or in the
bo_1012 slide code it reuses ([old N] = slide N of bo_1012.pptx).
Derived from make_bo_1012_deck.py (same helpers, badges, footer style).
Helpers: ../../pi_1008/deck_lib.py. After the build, an overflow check
estimates every text box and table bottom and asserts it is <= 7.0 in.
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent                      # quality_report/bo_1012/
sys.path.insert(0, str(HERE.parent.parent / "pi_1008"))     # shared deck_lib.py

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

import deck_lib as L
from deck_lib import (ACCENT, INK, MUTED, card, content_slide, figure, new_deck, notes,
                      section_slide, table, text_box)

AN = HERE.parent.parent                                      # campaign/analysis
ADMIN = Path("/Users/yilu/Documents/development/nd/research/admin/prompt/bogp")

RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x1E, 0x84, 0x49)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BOX = RGBColor(0xFD, 0xF6, 0xE3)       # pale gold for boxed key sentences

FOOT = "BO campaign + ERCOT v2 · 2026-10-13"
G1 = "Costs are g1-pending (option c adds +$28.188M to every design)"

# ---- sources (outline shorthand -> absolute path) ----------------------------
B = AN / "bo_stage2/BO_STAGE2.md"
B_ROUND = AN / "bo_stage2/ROUND_LOG.md"
R = AN / "random_baseline/RANDOM_BASELINE.md"
P30 = ADMIN / "30-stage2-bo-campaign.md"
P34 = ADMIN / "34-random-baseline.md"
P35 = ADMIN / "35-stage2-methods-arms.md"
P37 = ADMIN / "37-stop-rule-validation.md"
P38 = ADMIN / "38-rts-mechanism-checks.md"
PROG = ADMIN / "PROGRESS.md"
ML = ADMIN / "math-log.md"
LR = ADMIN / "living-report-multipem.md"
RVM = "methods review 10-13"
RVP = "power-systems review 10-13"

# ---- figure paths (edit here to swap) -----------------------------------------
F = {
    "front": AN / "bo_stage2/figs/fig1_front.png",
    "hv_traj": AN / "bo_stage2/figs/fig2_hv_trajectory.png",
    "calib": AN / "bo_stage2/figs/fig3_calibration.png",
    "front_designs": AN / "bo_stage2/figs/fig4_front_designs.png",
    "hv_band": AN / "random_baseline/figs/fig1_hv_band.png",
    "pct_hist": AN / "random_baseline/figs/fig2_percentile_hist.png",
    "s1": AN / "random_baseline/figs/fig3_s1_extremes.png",
    "grand": AN / "random_baseline/figs/fig4b_grand_front_unified.png",
}

n = 0


def nxt():
    global n
    n += 1
    return n


def bullets(slide, x, y, w, h, items, size=14, space_after=6):
    return text_box(slide, x, y, w, h, items, size=size, bullets=True, space_after=space_after)


def src(*paths):
    return "Sources: " + "; ".join(str(p) for p in paths)


def add_badges(slide, items):
    """Copied from pi_1008/v2/make_pi_1008_v2.py::add_badges (title shrinks to leave room)."""
    ph = slide.shapes.title
    if ph is not None:
        ph.left, ph.top, ph.width, ph.height = Inches(0.14), Inches(0.14), Inches(9.8), Inches(0.86)
        for p in ph.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(18 if len(ph.text_frame.text) > 60 else 21)
    y = Inches(0.16)
    for kind, text in items:
        color = RED if kind == "R" else GREEN
        label = ("ROADBLOCKER · " if kind == "R" else "DECISION · ") + text
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.05), y, Inches(3.15), Inches(0.36))
        shp.adjustments[0] = 0.3
        shp.fill.solid(); shp.fill.fore_color.rgb = color
        shp.line.fill.background(); shp.shadow.inherit = False
        tf = shp.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.01)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.name = L.FONT; r.font.size = Pt(8.5); r.font.bold = True; r.font.color.rgb = WHITE
        y += Inches(0.40)


def block_divider(prs, kicker, heading, sub=None, color=None):
    s = section_slide(prs, kicker, heading, sub)
    if color is not None:
        band = [sh for sh in s.shapes if sh.shape_type == 1][0]
        band.fill.solid(); band.fill.fore_color.rgb = color
    return s


def boxed(slide, x, y, w, h, lines, size=13, fill=BOX, line=L.ACCENT2):
    card(slide, x, y, w, h, fill=fill, line=line)
    return text_box(slide, x + Inches(0.15), y + Inches(0.05), w - Inches(0.3), h - Inches(0.1), lines,
                    size=size, anchor=MSO_ANCHOR.MIDDLE, space_after=3)


def colour_title(slide, color):
    for p in slide.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = color


def two_col(slide, y, left_head, left_items, right_head, right_items, h, size=14, space_after=8,
            left_line=None, right_line=None, head_colors=(ACCENT, ACCENT)):
    """Two 6.0-in columns (x = 0.5 and 6.8) in cards, each with a heading and bullets."""
    for x, head, items, ln, hc in [(0.5, left_head, left_items, left_line, head_colors[0]),
                                   (6.8, right_head, right_items, right_line, head_colors[1])]:
        card(slide, Inches(x), y, Inches(6.0), h, line=ln)
        text_box(slide, Inches(x + 0.15), y + Inches(0.08), Inches(5.7), Inches(0.42), head, size=16, bold=True,
                 color=hc)
        bullets(slide, Inches(x + 0.15), y + Inches(0.55), Inches(5.7), h - Inches(0.65), items, size=size,
                space_after=space_after)


prs = new_deck()

# =============================================================================
# 1  Title  (section_slide style + boxed "Three asks today")
# =============================================================================
s = section_slide(prs, "PI MEETING · 2026-10-13", "BO campaign + ERCOT v2: status and decisions, 2026-10-13")
nxt()
text_box(s, Inches(0.9), Inches(3.55), Inches(11.5), Inches(0.6),
         "Stage-2 q-ParEGO on RTS-GMLC (scenario C), the methods arms now running, and the first ERCOT "
         "results on corrected data.", size=15, color=RGBColor(0xE4, 0xE9, 0xF0))
boxed(s, Inches(0.9), Inches(4.3), Inches(11.5), Inches(1.8), [
    ("Three asks today", {"bold": True, "color": ACCENT, "size": 15}),
    "1.  Sign off the N1 gate amendment and the primary HV reference (slides 15–16).",
    "2.  Approve the redesigned ERCOT wave D (slide 22).",
    "3.  Decide: scenarios A/B before submission, or a scenario-C-only claim (slide 25).",
], size=14)
notes(s, "Deck rebuilt 10-13 from bo_1012 after two expert reviews (" + RVM + "; " + RVP + "). "
         "Built by make_bo_1013_deck.py from OUTLINE_1013.md (content source of truth).")

# =============================================================================
# 2  Claims table  [NEW]
# =============================================================================
s = content_slide(prs, "What each claim needs, and where it stands", nxt(), FOOT)
rows = [
    ["Claim", "Evidence (experiment)", "Status 10-13", "What would weaken it"],
    ["W  Workflow: pre-registered, auditable MOBO on a real PCM stops on its rule", "Stage-2 run (done)",
     "supported; \"unattended\" only up to qsub", "—"],
    ["S  The stop rule has a known operating characteristic", "37 (truth-known problems); 35 P1 r3–r4 (real data)",
     "running", "S0 acts as \"stop after 2 rounds\""],
    ["ii  Results are stable across replicates (start set + seed)", "35 S2, S3 vs P1", "running",
     "P2/P3 below 0.8 percentile, or recipe at base rate"],
    ["iii  Found designs follow the sensitivity structure (recipe)",
     "Stage-2 + random baseline; 35 P2/P3 out-of-sample", "in-sample only", "recipe fails on P2/P3"],
    ["A  The workflow does not depend on the acquisition", "35 N1 (qNEHVI, 1 seed); 37", "running",
     "N1 clearly better than q-ParEGO"],
    ["B  BO places above random; BO reaches the extremes", "random baseline (done)",
     "above at both refs, not robust; extremes clear", "—"],
    ["T  The workflow transfers to another grid", "ERCOT (36: screening + OAT only, no BO)", "no experiment", "—"],
    ["i  Convergence within budget", "—", "retired (not shown)", "—"],
]
table(s, Inches(0.5), Inches(1.35), Inches(12.3), rows, col_widths=[4.1, 3.2, 2.6, 2.4], size=12,
      row_h=Inches(0.48))
text_box(s, Inches(0.5), Inches(6.2), Inches(12.3), Inches(0.6),
         "Bottom line: Sobol continuations Q1–Q3 serve no claim (descriptive). Prompt 38 (mechanism) gates the "
         "shed axis for S, B and iii.", size=12, color=MUTED)
notes(s, "Claims W, S, ii, iii, A, B, T, i: " + RVM + "; P35 §5; P37. " + src(P35, P37, ML))

# =============================================================================
# 3  Three prompts on CRC  [old 31, updated]
# =============================================================================
s = content_slide(prs, "Three prompts run on CRC; prompt 38 is ready", nxt(), FOOT)
rows = [
    ["Work", "Question", "CRC work", "State 10-13", "Results"],
    ["Replicates + acquisition (35)",
     "Is q-ParEGO stable across replicates at a fixed budget of 64? Is one qNEHVI run consistent with it?",
     "7 arms: start sets S2, S3; Sobol continuations Q1–Q3; P1 rounds 3–4; qNEHVI N1",
     "running since 10-12 21:00; this is the last allowed relaunch",
     "first round ≈ 10-13; all ≈ 10-17 → 10-19"],
    ["ERCOT on corrected data (36)",
     "Re-run ERCOT for AIChE",
     "A base ×3, B screening 19, C selection 9",
     "A + B landed; base read done; C nuclear rows 6–11 rerun after a T0 fix",
     "C ≈ 10-14; wave D after redesign"],
    ["Stop-rule validation (37)",
     "Does the stop rule work on problems with a known truth?",
     "1,540 tasks, 0 PCM runs",
     "running since 10-12 evening",
     "≈ 10-13 → 10-15"],
    ["Mechanism checks (38)",
     "Why does shed fall as withheld MW rises? Why do 7 designs share 22.426 MWh?",
     "0 PCM runs (rsync of finished runs)",
     "ready to hand off",
     "≈ 2 days after hand-off"],
]
table(s, Inches(0.5), Inches(1.4), Inches(12.3), rows, col_widths=[1.9, 3.5, 2.9, 2.4, 1.6], size=12,
      row_h=Inches(0.95))
text_box(s, Inches(0.5), Inches(6.3), Inches(12.3), Inches(0.4),
         "CRC order: ERCOT first. Submission stays manual from the login node.", size=12, color=MUTED)
notes(s, "Sources: admin PROGRESS.md rows 1–10; LR 10-12u, 10-12v, 10-13a, 10-13b. 35 relaunch 8308211; "
         "37 CRC gate f5e4c14. Prompt 38 written 10-13.\n" + src(PROG, LR, P38))

# =============================================================================
# 4  Reproducibility card  [old 7]
# =============================================================================
s = content_slide(prs, "Reproducibility card: the campaign configuration (frozen)", nxt(), FOOT)
rows = [
    ["Item", "Setting"],
    ["Design space", "6 tiers (nuclear, pv, tail, wind_122, wind_303, wind_317) × 9 lattice levels ω ∈ {0.05 … 1.0} = 531,441 designs"],
    ["Scenario", "C: ρ = 2.0, B = $40/MWh"],
    ["Objectives (min)", "load_shed_mwh, true_curtailment_mwh, total_cost_less_synthetic_usd"],
    ["Starts", "32 Sobol points (shared by every comparison)"],
    ["GP training data", "the 32 starts only (10 backfill rows: report-only sensitivity)"],
    ["Acquisition", "q-ParEGO, q = 8: 8 Dirichlet weights, EI, kriging believer, discrete lattice candidates"],
    ["HV reference (fixed before round 1)", "nadir of the 32 starts + 10 % span; HV₀ = 6.606e17"],
    ["Stop rule", "2 consecutive rounds with gain < τ = 5.263e16 (8.0 % of HV₀; set by the shed materiality floor, conservative)"],
    ["Software", "bo-gp feat/qparego @ 60e6328; RTS runs are bit-reproducible"],
]
table(s, Inches(0.5), Inches(1.35), Inches(12.3), rows, col_widths=[3.3, 9.0], size=13, row_h=Inches(0.52))
notes(s, "Notes: B header; prompt 30.\n" + src(B, P30))

# =============================================================================
# 5  PEM model  [NEW]
# =============================================================================
s = content_slide(prs, "The \"PEM\" is a $40 bid on renewable output, not an electrolyzer load", nxt(), FOOT)
two_col(s, Inches(1.35), "What the model does", [
    "Renewable site: the unit splits. The parent keeps availability minus the PEM capacity. A _PEM unit gets "
    "the rest and offers it at B = $40/MWh. Withheld MWh count as hydrogen (20 kg/MWh).",
    "Nuclear site: p_min rises to p_max − C. The band is priced at B. The unit stays committed every hour.",
    "The PEM sells to the grid when the price is above B. It adds no load.",
], "What the model omits", [
    "No electrolyzer efficiency curve, minimum load, ramp limit or capex.",
    "No hydrogen revenue or storage.",
    "Cost objective = system cost minus the synthetic bid costs. This removes the nuclear unit's real fuel "
    "cost at any ω: a constant on RTS (the g1 item), a bias against the base case on ERCOT.",
], Inches(3.85), size=14, right_line=RED, head_colors=(ACCENT, RED))
boxed(s, Inches(0.5), Inches(5.4), Inches(12.3), Inches(1.1),
      [("Every result in this deck is about priced withholding of renewable and nuclear output. "
        "Electrolyzer physics is future work.", {"bold": True})], size=16)
notes(s, "Sources: parameters.py (renewable split, thermal patch); summarize_wave.py header note on synthetic cost; "
         + RVP + " §3.")

# =============================================================================
# 6  Chain flow + defect   [old 8]  [RED]
# =============================================================================
s = content_slide(prs, "The loop runs unattended except for qsub; CRC submission stays manual", nxt(), FOOT)
add_badges(s, [("R", "CRC compute nodes cannot qsub")])
steps = [
    ("array", "8 Prescient runs", ACCENT),
    ("collector", "objectives, push", ACCENT),
    ("acquisition", "grade, propose, push", ACCENT),
    ("submit next round", "fails: compute nodes are not submit hosts", RED),
]
bx, bw, bh, by0, bgap = Inches(0.9), Inches(4.2), Inches(0.95), Inches(1.45), Inches(0.42)
for i, (hd, sub, col) in enumerate(steps):
    y = by0 + (bh + bgap) * i
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y, bw, bh)
    shp.adjustments[0] = 0.15
    shp.fill.solid(); shp.fill.fore_color.rgb = L.CARD if col is ACCENT else RGBColor(0xFB, 0xE9, 0xE7)
    shp.line.color.rgb = col; shp.line.width = Pt(2 if col is RED else 1.25)
    shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    L._run(p, hd, 16, True, col)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    L._run(p2, sub, 12, False, INK)
    if i < len(steps) - 1:
        ax = bx + bw // 2
        con = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ax, y + bh, ax, y + bh + bgap)
        con.line.color.rgb = RED if i == len(steps) - 2 else ACCENT
        con.line.width = Pt(2.25)
        ln = con.line._get_or_add_ln()
        tail = ln.makeelement(L.qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
lx = Inches(0.55)
ytop = by0 + bh // 2
ybot = by0 + (bh + bgap) * 3 + bh // 2
for (x1, y1, x2, y2, arrow) in [(bx, ybot, lx, ybot, False), (lx, ybot, lx, ytop, False), (lx, ytop, bx, ytop, True)]:
    con = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    con.line.color.rgb = MUTED; con.line.width = Pt(1.5); con.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = con.line._get_or_add_ln()
        ln.append(ln.makeelement(L.qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
text_box(s, Inches(0.5), Inches(6.6), Inches(5.0), Inches(0.3), "Chain per round; dashed = next round.",
         size=10, color=MUTED)
card(s, Inches(5.7), Inches(1.4), Inches(7.1), Inches(5.45))
bullets(s, Inches(5.9), Inches(1.55), Inches(6.7), Inches(5.2), [
    "Unattended hops worked: both collectors, both acquisitions (grade → propose → commit → push).",
    ("The submit hop cannot work on CRC: compute nodes are not submit hosts.", {"bold": True, "color": RED}),
    "The WEDGED trap fired and left a clear record. Kay submitted round 2 by hand.",
    "r2 acquisition re-run on the Mac: weights and all 8 designs identical to CRC.",
    "Proposed fix (not implemented): submit the whole chain from the login node with -hold_jid; or a login-node cron watcher.",
    ("This limits claim W: \"unattended\" holds up to job submission.", {"bold": True}),
], size=16, space_after=12)
notes(s, "Notes: B pain log + audit table.\n" + src(B, B_ROUND))

# =============================================================================
# 7  HV trajectory + why gains are small  [old 9 + old 11]
# =============================================================================
s = content_slide(prs, "The stop rule fired at the earliest possible round: τ (8 %) is larger than BO's whole gain (6.9 %)",
                  nxt(), FOOT)
figure(s, F["hv_traj"], Inches(0.5), Inches(1.35), Inches(6.9), Inches(5.5), n)
card(s, Inches(7.6), Inches(1.4), Inches(5.2), Inches(5.45))
bullets(s, Inches(7.75), Inches(1.5), Inches(4.9), Inches(5.25), [
    "HV: 6.606e17 → 6.969e17 (+5.5 %) → 7.061e17 (+1.4 %). Both gains are under τ.",
    "5 of 16 BO designs lie outside the fixed box and earn zero HV. With a wider box the gains are 2.6 % and 2.0 %.",
    "The front is dense: 26 of 48 points are non-dominated.",
    ("The stop means \"gains below a conservative materiality bar\". It does not mean \"converged\".", {"bold": True}),
    "Prompt 37 tests this directly. Its pre-registered reading: if S0 matches \"stop after 2 rounds\", the paper "
    "says so and reports the τ at which the rule starts to adapt.",
], size=14, space_after=10)
notes(s, f"Figure: {F['hv_traj']}\nNotes: B §2 (old slides 9 and 11); P37 pre-registered reading.\n" + src(B, P37))

# =============================================================================
# 8  Final front + corners  [old 10]
# =============================================================================
s = content_slide(prs, "BO's front spans $510M with 29 GWh shed to zero shed at $582M", nxt(), G1)
figure(s, F["front"], Inches(0.5), Inches(1.3), Inches(12.3), Inches(3.8), n)
rows = [
    ["Best in", "Design(s)"],
    ["Shed", "4 designs tie at 0; cheapest r2#6 $582M; r1#5 = 0 shed + lowest curtailment at $668M"],
    ["Curtailment", "r1#5: 48 GWh, $668M (nuclear ω = 1.0)"],
    ["Cost", "r2#3: $510M (all tiers 0.05 except tail 1.0; 29 GWh shed, 1.26 TWh curtailed)"],
]
table(s, Inches(0.5), Inches(5.2), Inches(12.3), rows, col_widths=[1.8, 10.5], size=13, row_h=Inches(0.4))
notes(s, f"Figure: {F['front']}\nNotes: B §1. Footer: costs g1-pending (+$28.188M constant under option c).\n" + src(B))

# =============================================================================
# 9  Calibration   [old 12]  [RED]
# =============================================================================
s = content_slide(prs, "The calibration audit failed on 2 of 3 objectives and changed nothing", nxt(), FOOT)
add_badges(s, [("R", "corner extrapolation")])
figure(s, F["calib"], Inches(0.5), Inches(1.3), Inches(12.3), Inches(3.6), n)
rows = [
    ["Objective", "r1 RMS(z)", "r2 RMS(z)", "Pooled cov95", "Verdict"],
    ["shed", "0.33", "1.88", "94 %", "not calibrated"],
    ["curtailment", "0.55", "2.49", "88 %", "not calibrated"],
    ["cost", "0.91", "1.01", "100 %", "calibrated"],
]
table(s, Inches(0.5), Inches(5.0), Inches(6.9), rows, col_widths=[1.6, 1.2, 1.2, 1.4, 1.6], size=12, row_h=Inches(0.42))
bullets(s, Inches(7.6), Inches(4.95), Inches(5.2), Inches(1.95), [
    "r1 too wide, r2 overconfident at corners (r2#3 shed 29.2k vs 8.5k predicted).",
    "The audit had no decision role.",
    "Prompt 37 tests whether the audit predicts error.",
    ("Any restart: add corner points; consider shed as a constraint.", {"bold": True}),
], size=13, space_after=4)
notes(s, f"Figure: {F['calib']}\nNotes: B §3; P37.\n" + src(B, P37))

# =============================================================================
# 10  Random baseline: design + primary result  [old 14 + old 15]
# =============================================================================
s = content_slide(prs, "BO places at the 94th percentile of random 16-design batches; the lead is below τ", nxt(), FOOT)
figure(s, F["hv_band"], Inches(0.5), Inches(1.3), Inches(7.3), Inches(2.7), n, show_path=False)
figure(s, F["pct_hist"], Inches(7.9), Inches(1.3), Inches(4.9), Inches(2.7), n, show_path=False)
text_box(s, Inches(0.5), Inches(4.0), Inches(12.3), Inches(0.32),
         [(f"Left: {F['hv_band']}", {"color": L.PATH_GREY, "size": 8}),
          (f"Right: {F['pct_hist']}", {"color": L.PATH_GREY, "size": 8})], space_after=0)
text_box(s, Inches(0.5), Inches(4.35), Inches(12.3), Inches(0.35),
         "Design: 48 uniform random lattice designs; 100,000 random 16-subsets; \"above\" needs p̂ ≥ 0.90 and "
         "p_lo ≥ 0.50 at both refs (pre-registered).", size=12)
rows = [
    ["Ref", "BO HV/HV₀", "Random p10 / p50 / p90", "BO − median", "p̂48 [interval]", "Row"],
    ["R_frozen", "1.0688", "1.0158 / 1.0421 / 1.0651", "+0.027", "0.94 [0.84, 1.00]", "above"],
    ["R_union", "1.0460", "1.0083 / 1.0204 / 1.0311", "+0.026", "1.00 [1.00, 1.00]", "above"],
]
table(s, Inches(0.5), Inches(4.75), Inches(12.3), rows, col_widths=[1.4, 1.5, 3.2, 1.6, 2.4, 1.0], size=12,
      row_h=Inches(0.36))
boxed(s, Inches(0.5), Inches(5.95), Inches(12.3), Inches(0.9),
      "R_frozen (fixed before round 1) is the honest primary. R_union credits BO's corner picks; its 1.00 "
      "restates the extremes result.", size=14)
notes(s, f"Figures: {F['hv_band']}\n{F['pct_hist']}\nNotes: P34 §0.1, §3.1–3.2; R §2; statistics reviewer (expert review 10-12).\n" + "HOW TO READ THE TWO REFS. HV needs a reference point: the 'worst corner' of the box. A design that is worse than the ref on ANY objective earns zero HV. R_frozen = nadir of the 32 starts + 10 % span, fixed before round 1 (shed 15,102 MWh; curtailment 670,350 MWh; cost $652.0M). R_union = the same formula over all 96 designs (shed 32,100 MWh; curtailment 1,385,278 MWh; cost $684.3M); it equals BO_STAGE2's REF48. R_frozen leans toward random: 5 of BO's 16 picks are corner designs outside this box and earn zero. R_union leans toward BO: the bigger box credits those corners, so p-hat = 1.00 there mostly restates the S1 extremes result. Rule: 'above' only if above at BOTH refs. Lead with R_frozen. Each panel divides by its own HV0, so do not compare heights across the two panels (1.046 vs 1.069). Band = p10-p90 of random subsets (8-subsets at n = 40, 16-subsets at n = 48; 100,000 each)." + "\n" + src(R, P34))

# =============================================================================
# 11  Robustness   [old 16]  [RED]
# =============================================================================
s = content_slide(prs, "The percentile lead depends on load-shed differences smaller than the floor", nxt(), FOOT)
add_badges(s, [("R", "sensitivity to sub-floor shed differences")])
card(s, Inches(0.5), Inches(1.35), Inches(6.05), Inches(3.4))
text_box(s, Inches(0.7), Inches(1.45), Inches(5.65), Inches(3.2), [
    ("Pre-registered", {"bold": True, "color": ACCENT, "size": 16}),
    "1,000 perturbed draws (floors 3,000 MWh / 5,000 MWh / $0.5M).",
    "467 above, 531 comparable, 2 matched.",
    "0.467 < 0.8 → \"not robust to floor-level output perturbation\".",
    ("The flag stays.", {"bold": True}),
], size=16, space_after=10)
card(s, Inches(6.75), Inches(1.35), Inches(6.05), Inches(3.4), line=L.ACCENT2)
text_box(s, Inches(6.95), Inches(1.45), Inches(5.65), Inches(3.2), [
    ("POST-HOC (not pre-registered)", {"bold": True, "color": L.ACCENT2, "size": 16}),
    "The shed floor (3,000 MWh) is 20 % of the R_frozen shed width (15,102 MWh).",
    "Without shed noise, 0.89 of draws keep \"above\"; clipping shed at 0 gives 0.48.",
    "RTS runs are deterministic, so this is a materiality stress test, not measurement noise.",
], size=16, space_after=10)
boxed(s, Inches(0.5), Inches(4.95), Inches(12.3), Inches(1.9), [
    ("Reading", {"bold": True, "color": ACCENT}),
    "\"BO's percentile lead depends partly on load-shed differences smaller than 3,000 MWh.\" "
    "Also: p̂ = 0.94 sits only 0.04 above the 0.90 cutoff.",
    "Shed materiality is restated after prompt 38 (22.426 MWh plateau).",
], size=16)
notes(s, "Notes: R §3; strategy-session check 10-12; statistics reviewer (expert review 10-12); P38.\n" + src(R, P38))

# =============================================================================
# 12  Extremes  [old 17]
# =============================================================================
s = content_slide(prs, "BO's clearest result: both extremes lie beyond every random design", nxt(), G1)
figure(s, F["s1"], Inches(0.5), Inches(1.3), Inches(12.3), Inches(3.4), n)
rows = [
    ["Objective", "BO best", "Best of 48 random", "BO ahead by > floor?"],
    ["shed", "0", "0", "no (tie)"],
    ["curtailment", "48,444 MWh (r1#5)", "61,739 (rnd#15)", "yes, 13.3 GWh"],
    ["cost", "$510.1M (r2#3)", "$536.0M (rnd#42)", "yes, $25.9M"],
]
table(s, Inches(0.5), Inches(4.8), Inches(7.0), rows, col_widths=[1.3, 1.9, 1.8, 2.2], size=12, row_h=Inches(0.4))
bullets(s, Inches(7.7), Inches(4.75), Inches(5.1), Inches(2.15), [
    "No random 16-batch came within a floor of either extreme.",
    "Both extremes are corners.",
    "A minimum-build heuristic would likely find the cost corner; it was not run.",
    "The curtailment extreme puts the nuclear band at ω = 1.0: its hydrogen is nuclear energy, and gas replaces it.",
], size=12, space_after=4)
notes(s, f"Figure: {F['s1']}\nNotes: R §4.2; power-systems reviewer (expert review 10-12); " + RVP + ".\n" + src(R))

# =============================================================================
# 13  Design recipe   [old 19]  [GREEN]
# =============================================================================
PAPER = ("Every knee design, under either normalization, places the nuclear retrofit at its lattice minimum "
         "(ω = 0.05, 20 MW) and puts most of the PEM capacity (55–79%) on the two largest wind sites (buses 303 "
         "and 317). Total PEM capacity sets the position along the front, not membership of the knee region. "
         "Knees span 1.49–2.08 GW as the normalization changes. Over that range, cost rises about 5% while "
         "curtailment falls about 44%. The total therefore expresses a choice between curtailment and cost, "
         "not a recommended size.")
s = content_slide(prs, "The recipe narrows to two rules; prompt 35 tests them out of sample", nxt(), G1)
add_badges(s, [("G", "approve the narrowed recipe wording")])
rows = [
    ["Normalization", "Knee", "Source", "Nuclear ω", "Total MW", "wind_303+317 share", "wind_122 ω", "In old R"],
    ["full", "rnd#2", "random", "0.05", "1,491", "0.70", "0.169", "no"],
    ["full", "n0b#8", "start", "0.05", "1,699", "0.79", "0.169", "yes"],
    ["full", "r1#2", "BO", "0.05", "1,705", "0.68", "0.288", "yes"],
    ["clipped", "r2#1", "BO", "0.05", "1,825", "0.58", "0.525", "no"],
    ["clipped", "rnd#43", "random", "0.05", "2,084", "0.55", "0.644", "no"],
    ["clipped", "r1#2", "BO", "0.05", "1,705", "0.68", "0.288", "yes"],
]
table(s, Inches(0.5), Inches(1.3), Inches(12.3), rows, col_widths=[1.5, 1.1, 1.1, 1.2, 1.2, 2.0, 1.3, 1.1],
      size=11, row_h=Inches(0.28))
card(s, Inches(0.5), Inches(3.4), Inches(4.6), Inches(1.65))
text_box(s, Inches(0.65), Inches(3.45), Inches(4.3), Inches(1.55), [
    ("Holds", {"bold": True, "color": GREEN}),
    "Nuclear at 0.05; wind_303 + wind_317 share ≥ 0.5 (0.55–0.79).",
    ("Fails", {"bold": True, "color": RED}),
    "The 1.6–2.0 GW band; the wind_122 cap (drop it).",
], size=12, space_after=2)
boxed(s, Inches(5.3), Inches(3.4), Inches(7.5), Inches(1.65), [
    ("Key point", {"bold": True, "color": ACCENT}),
    "Total MW is the position along the front, not a size rule. Across 1.49 → 2.08 GW cost rises about 5 % "
    "and curtailment falls about 44 %. The size is a curtailment-vs-cost choice.",
    "Physics: a nuclear PEM diverts cheap baseload, so it never wins on cost.",
], size=12)
text_box(s, Inches(0.5), Inches(5.12), Inches(12.3), Inches(0.32), "Caveats", size=12, bold=True, color=MUTED)
text_box(s, Inches(0.5), Inches(5.42), Inches(12.3), Inches(1.45), [
    "Designs within 0.03 of the knee distance include some with wind share < 0.5, so the share rule describes "
    "the knees, not the whole front.",
    "ω = 0.05 is the lattice floor, so \"nuclear at the floor\" means \"no nuclear retrofit\" in practice.",
    "r2#1 was already outside the old R at write time (a gap in the prompt-34 pre-submission check).",
    "R′ was derived from P1; P2/P3 are the out-of-sample test.",
], size=12, bullets=True, space_after=2)
notes(s, "Notes: R §4.3 iii; power-systems reviewer (expert review 10-12); P35 (P2/P3 out-of-sample).\n"
         "Proposed paper wording (power-systems reviewer, verbatim):\n\"" + PAPER + "\"\n" + src(R, B, P35))

# =============================================================================
# 14  Claim scope  [old 20]
# =============================================================================
s = content_slide(prs, "Paper wording: what we can say now and what we cannot", nxt(), FOOT)
card(s, Inches(0.5), Inches(1.35), Inches(6.9), Inches(5.5), line=GREEN)
text_box(s, Inches(0.7), Inches(1.45), Inches(6.5), Inches(0.45), "Can say", size=18, bold=True, color=GREEN)
text_box(s, Inches(0.7), Inches(1.95), Inches(6.5), Inches(4.8), [
    "\"An automated, pre-registered MOBO campaign on a real PCM terminated under its rule within budget.\"",
    "\"This BO run's HV exceeded 94 % of random 16-design batches (+0.027 HV₀, below our 8 % materiality bar, "
    "sensitive to sub-floor load-shed differences). Its clearer distinction is reaching the cost and curtailment "
    "extremes, $26M and 13 GWh beyond any of 48 random designs.\"",
    "Found designs are consistent with the LOCATION sensitivity structure (claim iii: supported).",
    "The narrowed recipe (slide 13).",
], size=15, bullets=True, space_after=10)
card(s, Inches(7.6), Inches(1.35), Inches(5.2), Inches(5.5), line=RED)
text_box(s, Inches(7.8), Inches(1.45), Inches(4.8), Inches(0.45), "Cannot say", size=18, bold=True, color=RED)
text_box(s, Inches(7.8), Inches(1.95), Inches(4.8), Inches(4.8), [
    "\"beat random\", \"significant\", \"robust\", any p-value",
    "anything about q-ParEGO in general (one seed) or seed stability (claim ii: not shown)",
    "\"converged\" (claim i: not shown)",
    "anything vs Sobol or a heuristic; the R_union [1, 1] as certainty; the ~70 % model-world figure",
    "anything beyond scenario C (A/B not run)",
    "reliability benefit of PEM (prompt 38 open)",
], size=15, bullets=True, space_after=8)
notes(s, "Notes: B claim scope; ML §5.2; statistics reviewer (expert review 10-12); " + RVM + ".\n" + src(B, ML))

# =============================================================================
# 15  Deviations register  [NEW]
# =============================================================================
s = content_slide(prs, "Pre-registered or post hoc: every deviation on one page", nxt(), FOOT)
rows = [
    ["Item", "Pre-registered?", "What changed", "When", "Disclosure"],
    ["Stop rule, τ, R_frozen", "yes (prompt 30)", "—", "10-06", "—"],
    ["M = 3 objectives", "decided after the 32 starts landed", "—", "10-06", "state in paper"],
    ["Random-baseline wording", "yes (prompt 34)", "—", "10-10", "—"],
    ["Primary reference for 35 = R_union", "frozen 10-12, after R_union gave p̂ = 1.00",
     "lead reference changed from R_frozen", "10-12", "justify or make R_frozen co-primary (decision)"],
    ["Shed-noise diagnostic, recipe narrowing", "post hoc", "—", "10-12", "labelled post hoc"],
    ["N1 replay gate", "yes (two-sided)", "made one-sided after the result", "10-12", "P13 (next slide)"],
    ["37 Amendment 1", "before results", "400 local 6-D random runs quarantined", "10-12", "PREREG amendment"],
    ["ERCOT smoke gate", "amended before results", "cost ±15 % vs authors' SCUC", "10-12", "ERCOT_COST_FIX"],
    ["ERCOT wave-D site rule", "yes, text unchanged",
     "committed after B landed, before any per-site value was read", "10-13", "CARRIED_OVER / T3"],
    ["35 relaunch, software stack", "—", "~/.local packages are part of the frozen stack", "10-12", "P14"],
]
table(s, Inches(0.5), Inches(1.35), Inches(12.3), rows, col_widths=[3.0, 2.7, 3.2, 0.8, 2.6], size=11,
      row_h=Inches(0.44))
notes(s, "Sources: " + RVM + " §3.2–3.3; P30, P34, P35 §3.2, P37 Amendment 1, ERCOT_COST_FIX (df61ab4).")

# =============================================================================
# 16  N1 gate  [NEW]  [GREEN]
# =============================================================================
s = content_slide(prs, "The N1 gate was changed after its result: three items to sign off", nxt(), FOOT)
add_badges(s, [("G", "PI sign-off")])
two_col(s, Inches(1.35), "What happened", [
    "Pre-registered gate (35 §3.2): the qNEHVI replay median must lie INSIDE the q-ParEGO smoke range at "
    "n = 16 and 24.",
    "Result: n = 16: 0.9166, inside. n = 24: 0.9534, above the q-ParEGO maximum 0.9518.",
    ("The gate failed HIGH.", {"bold": True, "color": RED}),
    "Amendment 10-12, after the result: one-sided (median ≥ q-ParEGO minimum) → PASS. N1 now runs.",
], "Sign off", [
    "(a) Accept the amendment and its disclosure text (P13).",
    "(b) N1 vs P1 is reported as \"consistent with\" at most: one seed; model, acquisition and noise handling "
    "change together.",
    "(c) Fix now: if qNEHVI beats q-ParEGO on real runs by more than the P1–P3 spread, the paper reports both "
    "and keeps q-ParEGO as the workflow default because ______ (agree today: cost, simplicity, or prior "
    "pre-registration).",
], Inches(5.45), size=14, space_after=10, right_line=GREEN, head_colors=(ACCENT, GREEN))
notes(s, "Sources: P35 §3.2 AMENDMENT 10-12; replay_gate.json; " + RVM + ".\n" + src(P35))

# =============================================================================
# 17  Prompt 35 decision rule  [NEW]
# =============================================================================
s = content_slide(prs, "Prompt 35 decides claim (ii) by a fixed four-part rule", nxt(), FOOT)
bullets(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.45), [
    ("Claim (ii) holds only if ALL four parts pass:", {"bold": True}),
    ("(1) P1, P2, P3 each p̂48 ≥ 0.8 at R_union, raw and floor-robust HV.", {"level": 1}),
    ("(2) HV range across P1–P3 at n = 64 ≤ 1 τ̄.", {"level": 1}),
    ("(3) Each replicate's top-3 knees satisfy R′ above the base rate.", {"level": 1}),
    ("(4) Curtailment extreme beyond the pool's best by > floor in ≥ 2 of 3.", {"level": 1}),
    "Reading fixed now: all pass → supported at C. P1 passes, P2/P3 fail → \"one of three\"; the result depends "
    "on the start set. All fail → not shown. The workflow claim stands in every case.",
    ("Likely outcome (methods review): the AND-rule fails on at least one part. Plan the \"one of three\" "
     "wording now.", {"bold": True}),
    "P1 rounds 3–4 test the live stop on real data: any post-stop round > τ is an early stop.",
    "Sobol Q1–Q3: descriptive only.",
], size=15, space_after=8)
notes(s, "Sources: P35 §5; " + RVM + ".\n" + src(P35))

# =============================================================================
# 18  Prompt 37 confirmatory test  [NEW]
# =============================================================================
s = content_slide(prs, "Prompt 37 has one confirmatory test with fixed acceptance thresholds", nxt(), FOOT)
bullets(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.45), [
    "Primary cell: emulator sample paths E-path-1…5 × q-ParEGO × live rule S0 (τ = 1·τ_j, w = 2, q = 8, cap 4), "
    "50 seeds.",
    ("Accept only if: early-stop rate ≤ 20 % AND median regret ≤ 1 τ AND S0 on or below the fixed-budget line "
     "S-fix-k, on the primary problems AND on the needle problem.", {"bold": True}),
    "Comparators: S-fix-1…8, relative-gain rule, model-based rule (qNEHVI), OCD-style test, oracle.",
    "Pre-registered reading: S0 ≡ S-fix-2 almost everywhere → \"the rule acts as a 2-round budget at this τ\".",
    ("Limit: the emulators are fit on scenario-C data. They are not a second problem instance.",
     {"color": RED}),
], size=16, space_after=14)
notes(s, "Sources: P37 §4–5.\n" + src(P37))

# =============================================================================
# 19  RTS mechanism questions  [old 22 + old 27]
# =============================================================================
s = content_slide(prs, "Two RTS mechanism questions gate the shed axis; prompt 38 answers them", nxt(), FOOT)
rows = [
    ["Question", "Most likely answer", "Check (prompt 38)"],
    ["Why does shed fall as withheld MW rises (Spearman −0.94)?",
     "H1: in the RUC the slice is a $40 resource, so more thermal is committed and real-time headroom grows. "
     "H2: the slice sells when the price exceeds B, so it delivers in scarcity hours.",
     "Committed MW and headroom on shed hours, low- vs high-MW designs; shed hours vs price > B."],
    ["Is shed traded for reserve shortfall (H4)?", "Unknown",
     "Reserve shortfall vs total MW over 96 designs. If yes, the shed axis, τ and HV need rework."],
    ["Why do 7 front designs share shed = 22.426 MWh?", "One PEM-independent event",
     "Bus-hour of the shed; is it in the no-PEM base?"],
]
table(s, Inches(0.5), Inches(1.4), Inches(12.3), rows, col_widths=[2.6, 6.4, 3.3], size=12, row_h=Inches(0.9))
text_box(s, Inches(0.5), Inches(5.6), Inches(12.3), Inches(0.7),
         "Curtailment netting is closed: prompt 35 G1 closure PASS. Owner: hand off prompt 38 now (0 PCM runs).",
         size=13, color=MUTED)
notes(s, "Power-systems reviewer 10-12; " + RVP + "; parameters.py (renewable split), summarize_wave.py:113-116. "
         "Hypotheses until checked. 35 G1 closure: d876321.\n" + src(B, P38, P35))

# =============================================================================
# 20  ERCOT v2 base read  [old 28 + old 33]
# =============================================================================
s = content_slide(prs, "ERCOT v2 runs clean at the system level; the fuel mix needs a check", nxt(), FOOT)
rows = [
    ["Quantity (2019, base)", "v1 (mis-converted costs)", "v2 (corrected)"],
    ["Real-time LMP, mean", "$14.5/MWh", "$33.8/MWh"],
    ["Fixed cost", "$6.75 B", "$3.63 B"],
    ["Variable cost", "$1.19 B", "$4.46 B"],
    ["True curtailment", "9.44 M MWh", "5.92 M MWh"],
    ["Nuclear", "gen-1 cycles (≈ 25 starts)", "both units on all 8,760 h, 0 starts"],
    ["RUC quality", "TimeLimit races on capped days", "0 % capped; 1,460 / 1,460 OPTIMAL"],
    ["Replicate spread (total cost)", "±$80M bimodal", "0 (4 runs byte-identical)"],
    ["Coal / gas share", "—", "32 % / 34 % (2019 actual ≈ 20 % / 47 %, verify)"],
    ["PV energy", "—", "17.6 TWh from 8.3 GW (2019 actual ≈ 4–5 TWh, verify)"],
]
table(s, Inches(0.5), Inches(1.35), Inches(7.3), rows, col_widths=[2.2, 2.2, 2.9], size=12, row_h=Inches(0.42))
bullets(s, Inches(8.0), Inches(1.35), Inches(4.8), Inches(5.5), [
    "v1 used mis-converted cost data (start-up energy read as a heat rate). Every v1 output is replaced.",
    ("Corrected costs make nuclear baseload (informs NXT-D6).", {"bold": True}),
    ("The coal-gas split sets the marginal unit and the cost of replacing withheld wind. Check fuel mix and fuel "
     "prices against the TX-123BT paper and 2019 data before the AIChE figures.", {"color": RED}),
    "Deterministic is not exact: the RUC still stops at ≤ 1 % gap (median 0.72 %). A bid-0 placebo row gives "
    "the real ruler.",
], size=13, space_after=8)
notes(s, "T3 base read: ercot df61ab4 (ERCOT_COST_FIX.md); extract base_2019_v2 (e3082fa). v1 replicate spread: "
         "screening bimodality ±$79.5M (3edd682) and TL=360 probe |Δ| $56–127M (c9d28cd). v1 mapping error: old "
         "slide 28 (backup B6). Fuel mix and PV: " + RVP + ".\n" + src(LR, PROG))

# =============================================================================
# 21  ERCOT screening  [NEW]
# =============================================================================
s = content_slide(prs, "ERCOT screening: relief is real, but a third of the \"hydrogen\" is diverted wind", nxt(), FOOT)
cards21 = [
    ("0 MWh", "shed (4 base + 19 screening runs)"),
    ("$33.8/MWh", "mean RT LMP ($36.4 load-weighted)"),
    ("0.61–0.67", "MWh curtailment avoided per MWh H2 (top wind sites)"),
    ("0.38", "the same ratio at PV site 30"),
]
cw = Inches(2.95)
for i, (num, lab) in enumerate(cards21):
    cx = Inches(0.5) + (cw + Inches(0.17)) * i
    card(s, cx, Inches(1.35), cw, Inches(1.45))
    text_box(s, cx + Inches(0.15), Inches(1.4), cw - Inches(0.3), Inches(0.65), num, size=26, bold=True,
             color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    text_box(s, cx + Inches(0.15), Inches(2.05), cw - Inches(0.3), Inches(0.7), lab, size=12)
bullets(s, Inches(0.5), Inches(3.0), Inches(12.3), Inches(3.85), [
    "Shed is a checked constraint on ERCOT. Its span is zero, so any ERCOT BO uses two objectives: true "
    "curtailment and cost.",
    ("B = $40 lies above the mean price and above coal, so the PEM withholds in most hours. 33–40 % of the "
     "withheld energy would have served load; thermal replaces it.", {"bold": True}),
    "Screening row 10 shows the mechanism: 15.7 TWh withheld, 3.75 TWh of it curtailment; about 12 TWh replaced "
    "by thermal → +$354M.",
    "Report curtailment avoided per MWh of hydrogen next to relief per MW.",
], size=15, space_after=10)
notes(s, "Sources: ERCOT extracts screening_v2 (2ba803b) and base_2019_v2 (e3082fa); " + RVP + ".")

# =============================================================================
# 22  Wave D redesign  [NEW]  [GREEN]
# =============================================================================
s = content_slide(prs, "Wave D needs contrast, a placebo and a nuclear cost fix", nxt(), FOOT)
add_badges(s, [("G", "approve the redesigned wave D")])
two_col(s, Inches(1.35), "As proposed (rule 4fb3b1f, ruler = 0 → top 6)", [
    "Sites 163, 140, 146, 270, 110, 274 × 60/120/180 MW + nuclear × 3 = 21 rows.",
    ("Problems:", {"bold": True, "color": RED}),
    ("The six differ by 9 % in relief/MW and sit on 3 buses (120 ×3, 2 ×2, 75).", {"level": 1}),
    ("180 MW is 97 % of site 163's nameplate but 24 % of site 274's.", {"level": 1}),
    ("Nuclear rows carry the cost bias (slide 5).", {"level": 1}),
], "Redesign (21 rows)", [
    "163, 140, 146 (best per bus) + PV 116, PV 30 + the last-ranked wind site, each × 60/120/180 MW.",
    "+ one bid-0 placebo.",
    "+ nuclear × 2 after the cost fix.",
    "Plot the nameplate fraction too.",
], Inches(3.95), size=14, space_after=8, right_line=GREEN, head_colors=(ACCENT, GREEN))
boxed(s, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.15),
      "The pre-registered rule stays on record. The redesign is a disclosed amendment: the rule selects similar "
      "sites, so it cannot show the location effect.", size=15)
notes(s, "Sources: " + RVP + " §3; ercot 4fb3b1f (top-k rule). The 36 session builds after Kay approves.")

# =============================================================================
# 23  Reviewer requirements + status  [old 25 + status column]
# =============================================================================
s = content_slide(prs, "What reviewers will require, and what is now running", nxt(), FOOT)
rows = [
    ["#", "Requirement", "Why", "Cost", "Status 10-13"],
    ["1", "Stop-rule validation: early/late-stop rates and regret at stop on truth-known problems (replay, "
          "synthetic functions, GP emulator of Stage 2); compare with existing BO stop rules",
     "The stop rule is the most novel part of the workflow; today it has no operating characteristic",
     "0 PCM runs", "running (37)"],
    ["2", "Fixed-budget comparison: every arm to 32 + 4×8; apply the stop rule post hoc to the same trajectories",
     "BO stops on its own rule, baselines have none — budgets differ", "seed 1: +16 runs", "running (35)"],
    ["3", "q-ParEGO ≥ 3 replicates, ≥ 2 with NEW starts; qNEHVI ≥ 3 seeds paired by start set",
     "Acquisition-only seeds share 2/3 of the data and understate variance; qNEHVI is the comparator reviewers expect",
     "≈ 3–4k core-h at C", "partly running (2 new start sets, 1 qNEHVI seed)"],
    ["4", "Scenarios A and B as extra problem instances (cheapest: q-ParEGO + random)",
     "One instance = a case study, not a methods result", "≈ 1–5k core-h each", "no prompt (decision)"],
    ["5", "D3 noise floors + a floor-robust metric, pre-registered before new runs",
     "Today's differences sit below the floor", "spec first", "spec only"],
    ["6", "Emulator benchmark (≥ 20 seeds per method); calibration audit with a decision role",
     "Cheap method-level statistics; the audit failed 2 of 3 objectives with no consequence", "0 PCM runs",
     "running (37)"],
]
table(s, Inches(0.5), Inches(1.3), Inches(12.3), rows, col_widths=[0.35, 5.0, 3.6, 1.45, 1.9], size=11,
      row_h=Inches(0.5))
notes(s, "Contribution = the WORKFLOW (pre-registered, auditable, unattended MOBO with a materiality stop), for an applied-methods venue; not a new BO algorithm. "
         "Row 1 truth-known problems in full: replay A/B/C, synthetic functions with floor-matched noise, GP emulator of Stage 2. "
         "Downgraded: Sobol to one continuation per start set (descriptive); knee-only re-runs. Dropped: ≥ 20 real Stage-2 seeds, lower-τ continuation, 5-batch Sobol, other weather years, GP-fix restart. "
         "Total ≈ 4k (C + cheapest A/B) to ~14k core-h (full arms). Living report §9.7. Status column: PROGRESS 10-13.\n"
         + src(ML, LR, PROG))

# =============================================================================
# 24  AIChE plan  [old 26 rebuilt]
# =============================================================================
s = content_slide(prs, "AIChE: one message, five figures, a fallback", nxt(), FOOT)
boxed(s, Inches(0.5), Inches(1.35), Inches(12.3), Inches(1.25), [
    ("Message", {"bold": True, "color": ACCENT}),
    "\"An electrolyzer's grid value depends on where it sits: at congested wind buses it avoids about 0.6 MWh of "
    "curtailment per MWh of hydrogen; at solar or nuclear sites it avoids less at higher cost. Multi-objective BO "
    "maps these trade-offs in tens of production-cost runs.\"",
], size=14)
text_box(s, Inches(0.5), Inches(2.75), Inches(6.0), Inches(0.4), "Five figures", size=16, bold=True, color=ACCENT)
bullets(s, Inches(0.5), Inches(3.15), Inches(6.0), Inches(2.6), [
    "(1) PEM-model schematic (slide 5).",
    "(2) RTS 96-design front (backup B3).",
    "(3) RTS knee composition (slide 13).",
    "(4) ERCOT relief/MW and Δcurt/H2 for all 16 sites + nuclear.",
    "(5) ERCOT size response with the placebo band.",
], size=13, space_after=5)
text_box(s, Inches(6.8), Inches(2.75), Inches(6.0), Inches(0.4), "Calendar", size=16, bold=True, color=ACCENT)
rows = [
    ["Dates", "Work"],
    ["10-13 → 10-19", "35 lands; 37 lands; ERCOT C + fuel-mix check; wave D redesign submitted."],
    ["10-20 → 10-25", "Wave D lands; 38 done; figures drafted."],
    ["10-26 → 11-01", "Figures frozen 11-01."],
    ["11-02 → 11-08", "Talk."],
]
table(s, Inches(6.8), Inches(3.15), Inches(6.0), rows, col_widths=[1.5, 4.5], size=12, row_h=Inches(0.4))
boxed(s, Inches(0.5), Inches(5.85), Inches(12.3), Inches(0.95), [
    ("Fallback: if 35 is late, figures 2–3 use the existing single run with its stated limits. If wave D is late, "
     "figure 5 is dropped.", {"bold": False}),
], size=14)
notes(s, "Rebuilt from old slide 26 (execution order). Message and figure list: OUTLINE_1013 §24; " + RVP + ". "
         "Figure freeze 11-01, talk 11-02 → 11-08 as in bo_1012.\n" + src(LR, PROG))

# =============================================================================
# 25  Decisions today
# =============================================================================
s = content_slide(prs, "Decisions today", nxt(), FOOT)
colour_title(s, GREEN)
dec = [
    ("D1", "Sign off the N1 amendment (a, b, c on slide 16)."),
    ("D2", "Primary HV reference for 35: keep R_union with a written justification, or make R_frozen co-primary."),
    ("D3", "Scenarios A/B before submission, or a claim stated for scenario C only."),
    ("D4", "Approve the redesigned wave D (+ placebo; nuclear after the cost fix)."),
    ("D5", "ERCOT fuel-mix check before the AIChE figures (owner: 36 session)."),
    ("D6", "Hand off prompt 38 now."),
    ("D7", "g1 ruling (NXT-D7) still open; ERCOT nuclear rows have the same accounting issue."),
]
for i, (k, t) in enumerate(dec):
    y = Inches(1.4) + Inches(0.78) * i
    card(s, Inches(0.5), y, Inches(12.3), Inches(0.68))
    tag = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), y + Inches(0.11), Inches(0.9), Inches(0.46))
    tag.adjustments[0] = 0.3; tag.fill.solid(); tag.fill.fore_color.rgb = GREEN
    tag.line.fill.background(); tag.shadow.inherit = False
    tf = tag.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    L._run(p, k, 16, True, WHITE)
    text_box(s, Inches(1.75), y, Inches(10.9), Inches(0.68), t, size=15, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "D1/D2 of bo_1012 (methods design, BoTorch) were approved 10-12 (NXT-D9). D1 → slide 16; D2 → slides 15–16; "
         "D3 → slides 14, 23; D4 → slide 22; D5 → slide 20; D6 → slide 19; D7 → slide 5.")

# =============================================================================
# BACKUP
# =============================================================================
s = block_divider(prs, "BACKUP", "Backup", "B1–B7: supporting slides.")
nxt()
notes(s, "Backup slides B1–B7.")

# B1  Timeline  [old 5 + rows]
s = content_slide(prs, "B1 · Timeline 09-19 → 10-13", nxt(), FOOT)
rows = [
    ["Date", "Event"],
    ["09-19 → 09-22", "Stage-2 starts: 32 Sobol designs (n0 + n0b) run on CRC"],
    ["10-06", "M = 3 approved: load shed, true curtailment, cost less synthetic"],
    ["10-07", "Chain launched (block_1); 2 harness fixes first"],
    ["10-08", "Round 1 results; acquisition WEDGED at qsub (compute node is no submit host)"],
    ["10-09", "Round 2 submitted by hand from the login node (≈ 20 h lost)"],
    ["10-10", "Round 2 results; stop rule fires → BO_DONE; BO_STAGE2.md written and reviewed"],
    ["10-10", "Prompt 34 (random baseline) pre-registered, 3-agent review"],
    ["10-11", "48 random designs built, run (11:00 → 20:22), analysed (23:29)"],
    ["10-12", "Expert review of the random-baseline result (statistics + power systems)"],
    ["10-12", "Prompts 35, 36, 37 written and launched"],
    ["10-13", "ERCOT waves A + B landed; base read done"],
]
table(s, Inches(0.5), Inches(1.35), Inches(12.3), rows, col_widths=[2.0, 10.3], size=13, row_h=Inches(0.44))
notes(s, "Timeline from B (status line, pain log), ROUND_LOG.md, R, P34 and LR rows 10-06 … 10-13b; PROGRESS 10-13.\n"
         + src(B, B_ROUND, R, P34, LR, PROG))

# B2  Where BO went  [old 13]
s = content_slide(prs, "B2 · BO moves to the lattice edges; nuclear leaves the floor only at the reliability corners",
                  nxt(), FOOT)
figure(s, F["front_designs"], Inches(0.5), Inches(1.3), Inches(7.8), Inches(4.05), n)
rows = [
    ["Set", "Tiers at an edge (of 6)", "Fully interior"],
    ["starts", "0.75", "44 %"],
    ["r1", "1.9", "0 %"],
    ["r2", "3.25", "0 %"],
]
table(s, Inches(0.9), Inches(5.45), Inches(6.9), rows, col_widths=[1.4, 3.0, 2.5], size=12, row_h=Inches(0.34))
card(s, Inches(8.5), Inches(1.4), Inches(4.3), Inches(5.45))
bullets(s, Inches(8.65), Inches(1.55), Inches(4.0), Inches(5.2), [
    "Nuclear is bang-bang: 0.05 on the knees, 1.0 at the reliability corners. Spearman(nuclear MW, cost) = 0.89.",
    "wind_303 and wind_317 are loaded up on the front (mean ω 0.65 / 0.61).",
    "Tail goes to 1.0 in r2, but it correlates with nothing. Read it as EI pushing a flat direction, not a tail effect.",
], size=16, space_after=14)
notes(s, f"Figure: {F['front_designs']}\nNotes: B §4.\n" + src(B))

# B3  Grand front  [old 18]
s = content_slide(prs, "B3 · Grand front of 96 designs", nxt(), FOOT)
figure(s, F["grand"], Inches(0.5), Inches(1.3), Inches(12.3), Inches(3.9), n)
bullets(s, Inches(0.6), Inches(5.3), Inches(12.2), Inches(1.6), [
    "Same panels and markers as slide 8. Of the 13 starts on the 48-point front, 6 are now dominated by random designs; all 13 BO front designs stay.",
    "32 designs on the 96-point front: BO 13 of 16 (81 %), random 12 of 48 (25 %), starts 7 of 32.",
    "A random 16-batch puts 6.6 designs (p10–p90 5–8) on the front of starts ∪ batch. BO puts 13.",
    "No random design dominates any BO design (set coverage 0.00).",
], size=13, space_after=4)
notes(s, f"Figure: {F['grand']}\nNotes: R §4.3–4.4.\n" + src(R))

# B4  Status card 10-12  [old 4]
s = content_slide(prs, "B4 · Status card 10-12", nxt(), FOOT)
cards = [
    ("48", "designs evaluated by the campaign (32 Sobol starts + 2 BO rounds of 8)", 40),
    ("26", "designs on the final 3-objective front (13 starts, 6 r1, 7 r2)", 40),
    ("+5.5 % → +1.4 %", "HV gain per round, under the 8 % stop threshold → BO_DONE 10-10", 28),
    ("48", "uniform random designs run as the baseline (10-11, 9.4 h wall)", 40),
    ("0.94", "percentile of BO's HV among random 16-design batches (R_frozen)", 40),
    ("0.83", "rounds per day (one 20 h stall: CRC compute nodes cannot submit jobs)", 40),
]
cw, ch, gx, gy = Inches(3.95), Inches(2.5), Inches(0.225), Inches(0.25)
for i, (num, lab, sz) in enumerate(cards):
    cx = Inches(0.5) + (cw + gx) * (i % 3)
    cy = Inches(1.4) + (ch + gy) * (i // 3)
    card(s, cx, cy, cw, ch)
    text_box(s, cx + Inches(0.2), cy + Inches(0.15), cw - Inches(0.4), Inches(1.0), num, size=sz, bold=True,
             color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    text_box(s, cx + Inches(0.2), cy + Inches(1.25), cw - Inches(0.4), Inches(1.15), lab, size=14)
notes(s, "Notes: B status line, §2, pain log; R §2; LR 10-11c.\n" + src(B, R, LR))

# B5  ERCOT wave C  [bo_1012 slide 32 wave-C bullet + notes]
s = content_slide(prs, "B5 · ERCOT wave C: nuclear rows failed on the T0 output; fixed and rerun", nxt(), FOOT)
bullets(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.45), [
    ("Wave C: nuclear rows ω = 0.05 and 0.2875 failed at the first RUC build.", {"bold": True}),
    "The PEM patch raises p_min, but the T0 output stays at the original 729 MW.",
    "CRC log selection_t1_v2_array.o1533432.6: \"Generator 1 has less output at T0 than is feasible to ramp up to\". "
    "Same family as 1eaf05e (startup capacity).",
    "The thermal patch also sets ramp_up_60min = PEM capacity (ω × 2,430 MW), so the egret T0 check depends on ω.",
    ("Fix (7c5fb8d): T0 output ≥ new p_min. A build-only test reproduces the error without the fix and passes 6/6 "
     "with it.", {"bold": True}),
    "All nuclear rows 6–11 are rerun, so that the ω sweep shares one T0.",
    "Old rows 6–11 are quarantined in runs/quarantine_T0_729, not deleted.",
], size=15, space_after=10)
notes(s, "Wave C failure: 02b147d, CRC log selection_t1_v2_array.o1533432.6. Fix 7c5fb8d. From bo_1012 slide 32.\n"
         + src(LR, PROG))

# B6  ERCOT v1 cost data  [old 28]
s = content_slide(prs, "B6 · ERCOT v1: the cost data were mis-converted", nxt(), FOOT)
bullets(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.45), [
    ("gen.csv HR_avg_0 = Csu · Pmax / fuel price on all 128 thermal units (checked to machine precision). It is a start-up energy in MMBTU, read by Egret as a heat rate in BTU/kWh.", {"bold": True}),
    "HR_incr is a segment heat, read as a rate. Start-up heat = 1 MBTU and start cost = $1 on every unit.",
    "Effect: no-load $2.1–2.6M per hour on nuclear (intended ≈ $14k); start-ups nearly free; marginal costs scaled by ΔP/1000, so the merit order is inverted (CT $2.9/MWh vs $43 intended).",
    "gen-1 cycling, gen-93 never on, the ±$80M bimodality, TL=360, the rulers, screening, selection_t1: all artifacts or invalid.",
    "Fix (prompt 36): find or rebuild the conversion from the TX-123BT paper; corrected data in a NEW versioned directory; regression test against Egret's parsed costs.",
    ("AIChE promises ERCOT results: re-run base + 2 replicates + screening on the corrected data before the 11-01 figure freeze. Old ERCOT figures are not shown.", {"bold": True}),
], size=15, space_after=8)
notes(s, "Source: processing notebook gtep/123_bus_coal/ERCOT_BUS123_base_XC_editeddata/demo_processing_XC.ipynb. Correct mapping: HR_avg_0 = 1000·C(Pmin)/(fp·Pmin); "
         "HR_incr_k = 1000·(ΔC_k/ΔP_k)/fp; start cost = Csu·Pmax. Check GTEP 123-bus work from the same folder. Living report §9.8; TL360_VERDICT banner (ercot e844fa7).\n" + src(LR, ML))

# B7  Operations  [bo_1012 slide 34 minus the N1 bullet]
s = content_slide(prs, "B7 · Operations 10-12/13: env under /groups, import failure + relaunch, 37 fingerprint, CRC lessons",
                  nxt(), FOOT)
text_box(s, Inches(0.5), Inches(1.35), Inches(6.0), Inches(0.45), "35  Methods arms (scenario C)", size=17, bold=True,
         color=ACCENT)
bullets(s, Inches(0.5), Inches(1.85), Inches(6.0), Inches(3.4), [
    "G1 curtailment closure: PASS. G3: runs ≤ 9.4 h (random_C), BO rounds 16–22 h, acquisition 29 min.",
    "BoTorch env built on CRC under /groups (home quota is full).",
    "Launch 1 failed at import (PYTHONNOUSERSITE hid packages in ~/.local). Fixed; relaunched 21:00, the last allowed relaunch.",
], size=14, space_after=8)
text_box(s, Inches(6.8), Inches(1.35), Inches(6.0), Inches(0.45), "37  Stop-rule validation (truth known)", size=17,
         bold=True, color=ACCENT)
bullets(s, Inches(6.8), Inches(1.85), Inches(6.0), Inches(3.4), [
    "PREREG + Amendment 1: 400 local 6-D random runs quarantined; all 6-D work on CRC.",
    "Truth files: sha256 verified on CRC. Random-number fingerprint identical on Mac and CRC envs.",
    "1,540 tasks: q-ParEGO, qNEHVI and random on 2-D replay grids (method comparison) "
    "and 6-D emulator truths (stop-rule test).",
], size=14, space_after=8)
boxed(s, Inches(0.5), Inches(5.4), Inches(12.3), Inches(1.25), [
    ("CRC lessons (10-12/13)", {"bold": True, "color": ACCENT, "size": 14}),
    "Envs go under /groups. Test the real entry-path import, not a shallow import. "
    "A build-only check over the whole lattice catches patch bugs before a 7 h queue.",
], size=13)
notes(s, "35: d876321 (G1, G3), 1b17f42 (env), 6480264 + 8308211 (import fix, relaunch). "
         "37: 2d529ec (Amendment 1), f5e4c14 (CRC gate + fingerprint). N1 gate moved to slide 16.\n" + src(LR, PROG))


# =============================================================================
# overflow check (estimate only)
# =============================================================================
EMU_IN = 914400.0


def _est_text_h(tf, width_in):
    """Estimated rendered height (in) of a text frame of the given outer width."""
    inner = width_in - (tf.margin_left + tf.margin_right) / EMU_IN
    total = (tf.margin_top + tf.margin_bottom) / EMU_IN
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None] or [18.0]
        pt = max(sizes)
        cpl = max(1.0, inner * 72 / (0.5 * pt))
        lines = sum(max(1, math.ceil(len(seg) / cpl)) for seg in txt.replace("\v", "\n").split("\n"))
        sa = p.space_after.pt if p.space_after is not None else 0.0
        total += lines * 1.2 * pt / 72 + sa / 72
    return total


def overflow_check(prs):
    report, fails, spills = [], [], []
    for idx, sl in enumerate(prs.slides, 1):
        is_content = sl.shapes.title is not None
        bottom = 0.0
        for sh in sl.shapes:
            top, left = sh.top / EMU_IN, sh.left / EMU_IN
            w, h = sh.width / EMU_IN, sh.height / EMU_IN
            if sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                if top >= 7.0:                       # footer strip
                    continue
                est = _est_text_h(sh.text_frame, w)
                if sh.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE and est < h:
                    b = top + (h + est) / 2
                else:
                    b = top + est
                if est > h + 0.05:
                    spills.append((idx, sh.text_frame.text[:40], round(est, 2), round(h, 2)))
                if top < 1.0:                        # title and badges: check spill only
                    continue
                bottom = max(bottom, b)
                if is_content and (left < 0.5 - 1e-3 or left + w > 12.8 + 1e-3 or top < 1.3 - 1e-3):
                    fails.append((idx, "h/v range", sh.text_frame.text[:40]))
            elif sh.has_table:
                tot = 0.0
                cols = [c.width / EMU_IN for c in sh.table.columns]
                for r in sh.table.rows:
                    rh = max(_est_text_h(r.cells[j].text_frame, cols[j]) for j in range(len(cols)))
                    tot += max(rh, r.height / EMU_IN)
                b = top + tot
                bottom = max(bottom, b)
                if is_content and (left < 0.5 - 1e-3 or left + sum(cols) > 12.8 + 1e-3 or top < 1.3 - 1e-3):
                    fails.append((idx, "table range", ""))
            elif sh.shape_type == 13:                # picture
                if is_content and (top < 1.3 - 1e-3 or top + h > 6.9 + 1e-3 or left < 0.5 - 1e-3 or left + w > 12.8 + 1e-3):
                    fails.append((idx, "picture range", round(top + h, 2)))
        report.append((idx, round(bottom, 2)))
        if bottom > 7.0:
            fails.append((idx, "bottom", round(bottom, 2)))
    return report, fails, spills


# =============================================================================
# write
# =============================================================================
out = HERE / "bo_1013.pptx"
prs.save(out)

with open(HERE / "FIGURES_USED_1013.md", "w") as fh:
    fh.write("# Figures used in bo_1013.pptx (slide → path)\n\n")
    for sl, p in L.FIGURES_USED:
        flag = "" if Path(p).exists() else "  **MISSING**"
        fh.write(f"- slide {sl}: `{p}`{flag}\n")
    fh.write("- slide 6: chain flow drawn with shapes (no image)\n")
    if L.MISSING:
        fh.write("\n## Missing at build time\n\n" + "\n".join(f"- `{m}`" for m in L.MISSING) + "\n")

with open(HERE / "SLIDE_TEXT_1013.md", "w") as fh:
    fh.write("# bo_1013.pptx — slide text and speaker notes\n\n")
    for i, sl in enumerate(prs.slides, 1):
        fh.write(f"## Slide {i}\n\n")
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        fh.write(t + "\n")
                fh.write("\n")
            if sh.has_table:
                for r in sh.table.rows:
                    fh.write("| " + " | ".join(c.text_frame.text.replace("\n", " ") for c in r.cells) + " |\n")
                fh.write("\n")
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
            fh.write("**Notes:** " + sl.notes_slide.notes_text_frame.text.strip() + "\n\n")

report, fails, spills = overflow_check(prs)
print("wrote", out, "slides:", len(prs.slides), "missing figures:", len(L.MISSING))
print("overflow check (slide: estimated bottom, in):")
print("  " + "  ".join(f"{i}:{b}" for i, b in report))
for sp in spills:
    print("  SPILL (text taller than its box):", sp)
for f in fails:
    print("  FAIL:", f)
assert not [f for f in fails if f[1] == "bottom"], "overflow: estimated bottom > 7.0 in"
