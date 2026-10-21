"""Helper layer for building the PI-meeting deck with python-pptx.

Built ON the group-meeting template (`Sprint Week End June 6 2025.pptx`):
its master carries the ND logo (bottom-left) and a slide-number placeholder
(bottom-right). We load the template, drop its three example slides, and add
ours on its TITLE_ONLY / BLANK layouts, cloning the slide-number placeholder
onto every slide so the page number renders.

Conventions:
- 16:9 canvas, 13.333 x 7.5 in (same as the template); Arial throughout.
- One accent colour (ND navy) for titles/callouts, ND gold on dividers; greys for captions.
- Every figure slide prints its source path under the image AND in the
  speaker notes, so the author can swap figures by editing the path.
- A missing figure file is NOT fatal: a grey placeholder box with the path
  is drawn instead, and the path is recorded in MISSING for the report.
"""

from __future__ import annotations

import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

try:  # Pillow is used only to preserve image aspect ratios
    from PIL import Image
except Exception:  # pragma: no cover
    Image = None

TEMPLATE = Path("/Users/yilu/Documents/development/nd/research/admin/group_meeting/"
                "Sprint Week End June 6 2025.pptx")
LAYOUT_TITLE_ONLY = 4
LAYOUT_BLANK = 5

W, H = Inches(13.333), Inches(7.5)
FONT = "Arial"
ACCENT = RGBColor(0x0C, 0x23, 0x40)   # ND navy
ACCENT2 = RGBColor(0xAE, 0x93, 0x42)  # ND gold (sparingly)
INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
PATH_GREY = RGBColor(0x8A, 0x8A, 0x8A)
CARD = RGBColor(0xF2, 0xF4, 0xF7)
HEADER_BG = RGBColor(0xDC, 0xE3, 0xEE)
ALT_ROW = RGBColor(0xF6, 0xF6, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN = RGBColor(0xB2, 0x3A, 0x2E)

MISSING: list[str] = []
FIGURES_USED: list[tuple[int, str]] = []


def new_deck() -> Presentation:
    prs = Presentation(str(TEMPLATE))
    global W, H
    W, H = prs.slide_width, prs.slide_height  # 12192000 x 6858000 EMU (13.333 x 7.5 in)
    # drop the template's example slides (parts become unreachable and are not saved)
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)
    return prs


def _add_slide_number(slide):
    """Clone the layout's SLIDE_NUMBER placeholder and fill it with a slidenum field."""
    layout = slide.slide_layout
    for ph in layout.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes.clone_placeholder(ph)
            new = slide.placeholders[ph.placeholder_format.idx]
            new.text_frame.text = ""  # ensures a txBody exists on the clone
            txBody = new._element.txBody
            # replace paragraphs with one paragraph holding a slidenum field
            for p in list(txBody.findall(qn("a:p"))):
                txBody.remove(p)
            p = etree.SubElement(txBody, qn("a:p"))
            pPr = etree.SubElement(p, qn("a:pPr"))
            pPr.set("algn", "r")
            fld = etree.SubElement(p, qn("a:fld"))
            fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
            fld.set("type", "slidenum")
            rPr = etree.SubElement(fld, qn("a:rPr"))
            rPr.set("lang", "en-US")
            rPr.set("sz", "1000")
            latin = etree.SubElement(rPr, qn("a:latin"))
            latin.set("typeface", FONT)
            t = etree.SubElement(fld, qn("a:t"))
            t.text = "#"
            return new
    return None


def _blank(prs, layout_idx=LAYOUT_TITLE_ONLY):
    s = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    _add_slide_number(s)
    return s


def _run(p, text, size, bold=False, color=INK, italic=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    return r


def text_box(slide, x, y, w, h, lines, size=14, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False,
             space_after=4, line_spacing=1.05):
    """lines: str | list[str | (str, dict)] ; dict may set bold/color/size/level."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    _fill(tf, lines, size, color, bold, align, bullets, space_after, line_spacing)
    return tb


def _fill(tf, lines, size, color, bold, align, bullets, space_after, line_spacing):
    if isinstance(lines, str):
        lines = [lines]
    first = True
    for item in lines:
        opts = {}
        if isinstance(item, tuple):
            item, opts = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        level = opts.get("level", 0)
        prefix = ""
        if bullets and item:
            prefix = ("•  " if level == 0 else "–  ")
            if level:
                prefix = "      " + prefix
        _run(p, prefix + item, opts.get("size", size), opts.get("bold", bold),
             opts.get("color", color), opts.get("italic", False))


def title(slide, text, size=None):
    """Fill the layout's TITLE placeholder (template position: 0.14,0.14, 13.1 x 0.86 in)."""
    if size is None:
        size = 24 if len(text) <= 68 else 20  # two-line titles must fit the 0.86-in box
    ph = slide.shapes.title
    if ph is None:
        return text_box(slide, Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.9), text,
                        size=size, bold=True, color=ACCENT)
    tf = ph.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    _run(p, text, size, True, ACCENT)
    return ph


def footer(slide, text, n=None):
    """Small grey footer text, placed to the right of the ND logo (logo spans x < 1.5 in)."""
    if text:
        text_box(slide, Inches(1.65), Inches(7.08), Inches(8.3), Inches(0.3), text,
                 size=9, color=MUTED)
    # slide number comes from the cloned placeholder field; `n` kept for API compatibility


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def card(slide, x, y, w, h, fill=CARD, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.04
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def stat(slide, x, y, w, h, number, label, num_size=34, fill=CARD):
    card(slide, x, y, w, h, fill)
    text_box(slide, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), Inches(0.9),
             number, size=num_size, bold=True, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    text_box(slide, x + Inches(0.15), y + Inches(0.95), w - Inches(0.3), h - Inches(1.0),
             label, size=12, color=INK)


def figure(slide, path, x, y, w, h, slide_no=None, caption=None, show_path=True):
    """Place an image fitted inside (x,y,w,h), keep aspect ratio, print the path."""
    p = Path(path)
    FIGURES_USED.append((slide_no, str(p)))
    cap_h = Inches(0.0)
    if show_path or caption:
        cap_h = Inches(0.68 if (caption and show_path) else 0.32)
    box_h = h - cap_h
    if p.exists() and Image is not None:
        try:
            with Image.open(p) as im:
                iw, ih = im.size
            scale = min(w / iw, box_h / ih)
            pw, ph = int(iw * scale), int(ih * scale)
            px = x + (w - pw) // 2
            py = y + (box_h - ph) // 2
            slide.shapes.add_picture(str(p), px, py, pw, ph)
        except Exception:
            _placeholder(slide, x, y, w, box_h, p)
    else:
        MISSING.append(str(p))
        _placeholder(slide, x, y, w, box_h, p)
    lines = []
    if caption:
        lines.append((caption, {"color": INK, "size": 10}))
    if show_path:
        lines.append((str(p), {"color": PATH_GREY, "size": 8}))
    if lines:
        text_box(slide, x, y + box_h + Inches(0.02), w, cap_h, lines, size=10,
                 space_after=0)


def _placeholder(slide, x, y, w, h, p: Path):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
    shp.line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    _run(para, "FIGURE PLACEHOLDER\n" + p.name + "\n(file not found at build time)", 11,
         color=MUTED)


def table(slide, x, y, w, rows, col_widths=None, size=11, header=True,
          row_h=Inches(0.36)):
    nrows, ncols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrows, ncols, x, y, w, row_h * nrows)
    tbl = shp.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = int(w * cw / total)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            if header and r == 0:
                cell.fill.fore_color.rgb = HEADER_BG
            else:
                cell.fill.fore_color.rgb = ALT_ROW if r % 2 == 0 else WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            _run(para, str(val), size, bold=(header and r == 0), color=INK)
    return shp


def section_slide(prs, kicker, heading, sub=None, n=None):
    """Divider on the BLANK layout: navy card leaves the logo/page-number strip visible."""
    s = _blank(prs, LAYOUT_BLANK)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(1.2), W, Inches(5.1))
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()
    band.shadow.inherit = False
    text_box(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.5), kicker,
             size=14, color=ACCENT2, bold=True)
    text_box(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.6), heading,
             size=38, color=WHITE, bold=True)
    if sub:
        text_box(s, Inches(0.9), Inches(3.9), Inches(11.5), Inches(2.2), sub,
                 size=15, color=RGBColor(0xE4, 0xE9, 0xF0))
    return s


def content_slide(prs, heading, n=None, foot=None):
    s = _blank(prs, LAYOUT_TITLE_ONLY)
    title(s, heading)
    footer(s, foot or "", n)
    return s
